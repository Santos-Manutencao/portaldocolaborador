import os
import re
import sqlite3
import datetime
import openpyxl
import unicodedata
from difflib import SequenceMatcher

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PARENT_DIR = os.path.dirname(BASE_DIR)

def resolve_path(rel_name):
    p1 = os.path.join(BASE_DIR, rel_name)
    if os.path.exists(p1):
        return p1
    p2 = os.path.join(PARENT_DIR, rel_name)
    if os.path.exists(p2):
        return p2
    return p1

DB_PATH = os.path.join(BASE_DIR, 'sgp_database.db')
EXCEL_PATH = resolve_path('FUNCIONARIOS.xlsx')
FUNCAO_SALARIO_PATH = resolve_path('Funcao-Salario.xlsx')
DADOS_BANCARIOS_PATH = resolve_path('Dados_Bancarios_Colaboradores_Completo.xlsx')
TELEFONES_PATH = resolve_path('FUNCIONARIOS_TELEFONES.xlsx')
FOTOS_DIR = resolve_path('FOTOS CRACHAS')
ATESTADOS_DIR = resolve_path('ATESTADOS')

def clean_str(val):
    if val is None:
        return ''
    s = str(val).strip()
    return '' if s == 'None' else s

def clean_text_norm(t):
    if not t:
        return ''
    n = unicodedata.normalize('NFKD', str(t)).encode('ASCII', 'ignore').decode('ASCII')
    n = re.sub(r'[^A-Z0-9 ]', ' ', n.upper())
    return re.sub(r'\s+', ' ', n).strip()

def clean_date(val):
    if not val:
        return ''
    if isinstance(val, (datetime.datetime, datetime.date)):
        return val.strftime('%Y-%m-%d')
    s = str(val).strip()
    for fmt in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d', '%d/%m/%Y', '%d/%m/%y'):
        try:
            return datetime.datetime.strptime(s, fmt).strftime('%Y-%m-%d')
        except ValueError:
            pass
    return s

def clean_number(val):
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace('R$', '').replace(' ', '')
    if ',' in s and '.' in s:
        s = s.replace('.', '').replace(',', '.')
    elif ',' in s:
        s = s.replace(',', '.')
    try:
        return float(s)
    except ValueError:
        return 0.0

def normalize_name(name):
    if not name:
        return ''
    name = name.upper().strip()
    name = re.sub(r'\s+', ' ', name)
    return name

def init_db(conn):
    cursor = conn.cursor()
    cursor.execute('DROP TABLE IF EXISTS funcionarios;')
    cursor.execute('''
    CREATE TABLE IF NOT EXISTS funcionarios (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        matricula TEXT,
        nome TEXT NOT NULL,
        cpf TEXT,
        rg TEXT,
        data_nascimento TEXT,
        cargo TEXT,
        admissao TEXT,
        demissao TEXT,
        status TEXT DEFAULT 'Ativo',
        valor_hora REAL DEFAULT 0,
        salario_funcao REAL DEFAULT 0,
        premio_assiduidade REAL DEFAULT 0,
        salario_bruto REAL DEFAULT 0,
        salario REAL DEFAULT 0,
        local_trabalho TEXT DEFAULT 'Geral',
        rua TEXT,
        numero TEXT,
        bairro TEXT,
        cidade TEXT,
        cep TEXT,
        telefone TEXT,
        banco TEXT,
        agencia TEXT,
        conta TEXT,
        pix TEXT,
        foto_path TEXT,
        observacoes TEXT,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    ''')

    cursor.execute('''
    CREATE TABLE IF NOT EXISTS atestados (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        funcionario_id INTEGER NOT NULL,
        data_inicio TEXT NOT NULL,
        dias_afastamento INTEGER DEFAULT 1,
        data_retorno TEXT,
        medico_crm TEXT,
        cid TEXT,
        motivo TEXT,
        anexo_path TEXT,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY (funcionario_id) REFERENCES funcionarios(id) ON DELETE CASCADE
    );
    ''')

    cursor.execute('''
    CREATE TABLE IF NOT EXISTS horas_extras (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        funcionario_id INTEGER NOT NULL,
        data TEXT NOT NULL,
        percentual TEXT DEFAULT '50%',
        quantidade_horas REAL NOT NULL,
        motivo TEXT,
        valor_calculado REAL DEFAULT 0,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY (funcionario_id) REFERENCES funcionarios(id) ON DELETE CASCADE
    );
    ''')

    cursor.execute('''
    CREATE TABLE IF NOT EXISTS epis (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        funcionario_id INTEGER NOT NULL,
        ca TEXT,
        quantidade INTEGER DEFAULT 1,
        descricao TEXT NOT NULL,
        data_entrega TEXT NOT NULL,
        data_devolucao TEXT,
        observacao TEXT,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY (funcionario_id) REFERENCES funcionarios(id) ON DELETE CASCADE
    );
    ''')

    cursor.execute('''
    CREATE TABLE IF NOT EXISTS ferias (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        funcionario_id INTEGER NOT NULL,
        periodo_aquisitivo_inicio TEXT,
        periodo_aquisitivo_fim TEXT,
        periodo_concessivo_fim TEXT,
        data_inicio TEXT NOT NULL,
        dias INTEGER NOT NULL DEFAULT 30,
        data_retorno TEXT,
        abono_pecuniario INTEGER DEFAULT 0,
        dias_abono INTEGER DEFAULT 0,
        adiantamento_13 INTEGER DEFAULT 0,
        status TEXT DEFAULT 'Agendada',
        observacoes TEXT,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY (funcionario_id) REFERENCES funcionarios(id) ON DELETE CASCADE
    );
    ''')

    conn.commit()

def load_funcoes_salarios():
    roles_map = {}
    if not os.path.exists(FUNCAO_SALARIO_PATH):
        return roles_map
    try:
        wb = openpyxl.load_workbook(FUNCAO_SALARIO_PATH, data_only=True)
        ws = wb.active
        for r in range(3, ws.max_row + 1):
            cargo_raw = ws.cell(r, 1).value
            if not cargo_raw:
                continue
            cargo_norm = clean_text_norm(cargo_raw)
            sal = clean_number(ws.cell(r, 2).value)
            prem = clean_number(ws.cell(r, 3).value)
            peric = clean_number(ws.cell(r, 4).value)
            insal = clean_number(ws.cell(r, 5).value)
            roles_map[cargo_norm] = {
                'salario_funcao': sal,
                'premio_assiduidade': prem,
                'periculosidade': peric,
                'insalubridade': insal
            }
    except Exception as e:
        print(f"Erro ao carregar Funcao-Salario: {e}")
    return roles_map

def load_gastos_bruto(wb):
    gastos_map = {}
    sheet = None
    for name in wb.sheetnames:
        if 'GASTOS' in name.upper():
            sheet = wb[name]
            break
    if not sheet:
        return gastos_map

    for r in range(3, sheet.max_row + 1):
        mat = clean_str(sheet.cell(r, 1).value)
        nome = clean_text_norm(sheet.cell(r, 2).value)
        bruto = clean_number(sheet.cell(r, 4).value)
        if bruto > 0:
            if mat:
                gastos_map['mat_' + mat] = bruto
            if nome:
                gastos_map['nome_' + nome] = bruto
    return gastos_map

def get_photos_map():
    photos = {}
    if not os.path.exists(FOTOS_DIR):
        return photos
    for root, dirs, files in os.walk(FOTOS_DIR):
        for f in files:
            if f.lower().endswith(('.jpg', '.jpeg', '.png')):
                base_name = os.path.splitext(f)[0].strip().upper()
                rel_path = os.path.relpath(os.path.join(root, f), BASE_DIR).replace('\\', '/')
                photos[base_name] = rel_path
    return photos

def find_best_photo(name, photos_map):
    norm_name = normalize_name(name)
    if norm_name in photos_map:
        return photos_map[norm_name]
    
    best_score = 0
    best_file = None
    for p_name, p_path in photos_map.items():
        score = SequenceMatcher(None, norm_name, p_name).ratio()
        if score > 0.82 and score > best_score:
            best_score = score
            best_file = p_path
    return best_file

def load_telefones():
    telefones_map = {}
    if not os.path.exists(TELEFONES_PATH):
        return telefones_map
    try:
        wb = openpyxl.load_workbook(TELEFONES_PATH, data_only=True)
        for sname in wb.sheetnames:
            ws = wb[sname]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or not any(row):
                    continue
                mat = clean_str(row[0])
                nome = normalize_name(clean_str(row[1]))
                tel = ''
                for cell in row[2:]:
                    c_str = clean_str(cell)
                    if re.match(r'^\d{3}\.\d{3}\.\d{3}', c_str):
                        continue
                    if any(char.isdigit() for char in c_str) and len(c_str) >= 8:
                        tel = c_str
                        break
                if tel:
                    if mat:
                        telefones_map['mat_' + mat] = tel
                    if nome:
                        telefones_map['nome_' + nome] = tel
    except Exception as e:
        print(f"Erro ao carregar telefones: {e}")
    return telefones_map

def load_dados_bancarios():
    banco_map = {}
    if not os.path.exists(DADOS_BANCARIOS_PATH):
        return banco_map
    try:
        wb = openpyxl.load_workbook(DADOS_BANCARIOS_PATH, data_only=True)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not any(row):
                continue
            nome = normalize_name(clean_str(row[0]))
            banco = clean_str(row[1]) if len(row) > 1 else ''
            agencia = clean_str(row[2]) if len(row) > 2 else ''
            conta = clean_str(row[3]) if len(row) > 3 else ''
            pix = clean_str(row[4]) if len(row) > 4 else ''
            if nome:
                banco_map[nome] = {
                    'banco': banco,
                    'agencia': agencia,
                    'conta': conta,
                    'pix': pix
                }
    except Exception as e:
        print(f"Erro ao carregar dados bancários: {e}")
    return banco_map

def load_enderecos(wb):
    enderecos_map = {}
    sheet = None
    for name in wb.sheetnames:
        if 'ENDER' in name.upper():
            sheet = wb[name]
            break
    if not sheet:
        return enderecos_map

    for r in range(3, sheet.max_row + 1):
        mat = clean_str(sheet.cell(r, 1).value)
        nome = normalize_name(clean_str(sheet.cell(r, 2).value))
        rua = clean_str(sheet.cell(r, 3).value)
        num = clean_str(sheet.cell(r, 4).value)
        bairro = clean_str(sheet.cell(r, 5).value)
        cidade = clean_str(sheet.cell(r, 6).value)
        cep = clean_str(sheet.cell(r, 7).value)

        if re.match(r'^\d{3}\.\d{3}\.\d{3}', rua) or re.match(r'^\d{3}\.\d{3}\.\d{3}', num):
            continue

        addr_data = {
            'rua': rua,
            'numero': num,
            'bairro': bairro,
            'cidade': cidade,
            'cep': cep
        }
        if any(addr_data.values()):
            if mat:
                enderecos_map['mat_' + mat] = addr_data
            if nome:
                enderecos_map['nome_' + nome] = addr_data
    return enderecos_map

def run_import():
    print("Iniciando importação de dados para SGP SANTOS MANUTENÇÃO...")
    conn = sqlite3.connect(DB_PATH)
    init_db(conn)

    roles_map = load_funcoes_salarios()
    print(f"Tabela Funcao-Salario carregada: {len(roles_map)} funções mapeadas com assiduidade")

    photos_map = get_photos_map()
    print(f"Fotos encontradas: {len(photos_map)}")

    telefones_map = load_telefones()
    print(f"Telefones indexados: {len(telefones_map)}")

    banco_map = load_dados_bancarios()
    print(f"Dados bancários indexados: {len(banco_map)}")

    if not os.path.exists(EXCEL_PATH):
        print(f"Arquivo {EXCEL_PATH} não encontrado!")
        conn.close()
        return

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    enderecos_map = load_enderecos(wb)
    print(f"Endereços indexados: {len(enderecos_map)}")

    gastos_map = load_gastos_bruto(wb)
    print(f"Salários brutos da folha GASTOS indexados: {len(gastos_map)}")

    cursor = conn.cursor()

    cursor.execute("DELETE FROM atestados;")
    cursor.execute("DELETE FROM horas_extras;")
    cursor.execute("DELETE FROM epis;")
    cursor.execute("DELETE FROM funcionarios;")
    cursor.execute("DELETE FROM sqlite_sequence;")
    conn.commit()

    total_ativos = 0
    total_desligados = 0

    # 1. Import ATIVOS
    sheet_ativos = None
    for s in wb.sheetnames:
        if 'ATIV' in s.upper():
            sheet_ativos = wb[s]
            break

    if sheet_ativos:
        for r in range(3, sheet_ativos.max_row + 1):
            mat = clean_str(sheet_ativos.cell(r, 1).value)
            nome = clean_str(sheet_ativos.cell(r, 2).value)
            if not nome:
                continue

            cpf = clean_str(sheet_ativos.cell(r, 3).value)
            rg = clean_str(sheet_ativos.cell(r, 4).value)
            dt_nasc = clean_date(sheet_ativos.cell(r, 5).value)
            cargo = clean_str(sheet_ativos.cell(r, 6).value)
            admissao = clean_date(sheet_ativos.cell(r, 7).value)
            v_hora = clean_number(sheet_ativos.cell(r, 8).value)
            salario_funcao = clean_number(sheet_ativos.cell(r, 9).value)
            # Lê RIGOROSAMENTE a Coluna 10 (PREMIO) da planilha ATIVOS
            premio_assiduidade = clean_number(sheet_ativos.cell(r, 10).value)
            # Lê a Coluna 11 (LOCAL)
            local = clean_str(sheet_ativos.cell(r, 11).value)
            if not local or local.upper() == 'DESLIGADO':
                local = 'CDA'

            norm_nome = normalize_name(nome)
            clean_cargo = clean_text_norm(cargo)

            # Se salário da função não estiver preenchido, busca da tabela de funções
            if salario_funcao <= 0:
                r_info = roles_map.get(clean_cargo, {})
                salario_funcao = r_info.get('salario_funcao', 0.0)

            # Salário Bruto é RIGOROSAMENTE a somatória do salário da função mais o prêmio de assiduidade
            salario_bruto = round(salario_funcao + premio_assiduidade, 2)

            if v_hora:
                v_hora = round(v_hora, 2)
            elif salario_funcao > 0:
                v_hora = round(salario_funcao / 220, 2)
            else:
                v_hora = 0.0

            # Match address
            addr = enderecos_map.get('mat_' + mat) or enderecos_map.get('nome_' + norm_nome) or {}
            rua = addr.get('rua', '')
            num = addr.get('numero', '')
            bairro = addr.get('bairro', '')
            cidade = addr.get('cidade', '')
            cep = addr.get('cep', '')

            # Match phone
            tel = telefones_map.get('mat_' + mat) or telefones_map.get('nome_' + norm_nome) or ''
            if re.match(r'^\d{3}\.\d{3}\.\d{3}', tel):
                tel = ''

            # Match bank
            b_info = banco_map.get(norm_nome, {})
            banco = b_info.get('banco', '')
            agencia = b_info.get('agencia', '')
            conta = b_info.get('conta', '')
            pix = b_info.get('pix', '')

            # Match photo
            foto = find_best_photo(nome, photos_map)

            cursor.execute('''
            INSERT INTO funcionarios (
                matricula, nome, cpf, rg, data_nascimento, cargo, admissao,
                status, valor_hora, salario_funcao, premio_assiduidade, salario_bruto, salario, local_trabalho,
                rua, numero, bairro, cidade, cep, telefone,
                banco, agencia, conta, pix, foto_path
            ) VALUES (?, ?, ?, ?, ?, ?, ?, 'Ativo', ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ''', (
                mat, nome, cpf, rg, dt_nasc, cargo, admissao,
                v_hora, salario_funcao, premio_assiduidade, salario_bruto, salario_bruto, local,
                rua, num, bairro, cidade, cep, tel,
                banco, agencia, conta, pix, foto
            ))
            total_ativos += 1

    # 2. Import DESLIGADOS
    sheet_desl = None
    for s in wb.sheetnames:
        if 'DESL' in s.upper():
            sheet_desl = wb[s]
            break

    if sheet_desl:
        for r in range(3, sheet_desl.max_row + 1):
            mat = clean_str(sheet_desl.cell(r, 1).value)
            nome = clean_str(sheet_desl.cell(r, 2).value)
            if not nome:
                continue

            cpf = clean_str(sheet_desl.cell(r, 3).value)
            rg = clean_str(sheet_desl.cell(r, 4).value)
            cargo = clean_str(sheet_desl.cell(r, 5).value)
            admissao = clean_date(sheet_desl.cell(r, 6).value)
            demissao = clean_date(sheet_desl.cell(r, 7).value)
            salario_funcao = clean_number(sheet_desl.cell(r, 8).value)
            clean_cargo = clean_text_norm(cargo)
            r_info = roles_map.get(clean_cargo, {})
            premio_assiduidade = r_info.get('premio_assiduidade', 0.0)
            salario_bruto = round(salario_funcao + premio_assiduidade, 2)
            v_hora = round(salario_funcao / 220, 2) if salario_funcao else 0.0

            norm_nome = normalize_name(nome)
            addr = enderecos_map.get('mat_' + mat) or enderecos_map.get('nome_' + norm_nome) or {}
            tel = telefones_map.get('mat_' + mat) or telefones_map.get('nome_' + norm_nome) or ''
            if re.match(r'^\d{3}\.\d{3}\.\d{3}', tel):
                tel = ''
            b_info = banco_map.get(norm_nome, {})
            foto = find_best_photo(nome, photos_map)

            cursor.execute('''
            INSERT INTO funcionarios (
                matricula, nome, cpf, rg, cargo, admissao, demissao,
                status, valor_hora, salario_funcao, premio_assiduidade, salario_bruto, salario, local_trabalho,
                rua, numero, bairro, cidade, cep, telefone,
                banco, agencia, conta, pix, foto_path
            ) VALUES (?, ?, ?, ?, ?, ?, ?, 'Desligado', ?, ?, ?, ?, ?, '', ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ''', (
                mat, nome, cpf, rg, cargo, admissao, demissao,
                v_hora, salario_funcao, premio_assiduidade, salario_bruto, salario_bruto,
                addr.get('rua', ''), addr.get('numero', ''), addr.get('bairro', ''), addr.get('cidade', ''), addr.get('cep', ''),
                tel,
                b_info.get('banco', ''), b_info.get('agencia', ''), b_info.get('conta', ''), b_info.get('pix', ''),
                foto
            ))
            total_desligados += 1

    # 3. Seed initial EPIs
    sheet_epi = None
    for s in wb.sheetnames:
        if 'EPI' in s.upper():
            sheet_epi = wb[s]
            break
    
    if sheet_epi:
        for r in range(6, sheet_epi.max_row + 1):
            ca = clean_str(sheet_epi.cell(r, 1).value)
            qtd = int(clean_number(sheet_epi.cell(r, 2).value) or 1)
            desc = clean_str(sheet_epi.cell(r, 3).value)
            dt = clean_date(sheet_epi.cell(r, 4).value) or datetime.date.today().strftime('%Y-%m-%d')
            if desc:
                cursor.execute("SELECT id FROM funcionarios WHERE cargo LIKE '%MANTENEDOR%' AND status = 'Ativo' LIMIT 15")
                emp_ids = cursor.fetchall()
                for (eid,) in emp_ids:
                    cursor.execute('''
                    INSERT INTO epis (funcionario_id, ca, quantidade, descricao, data_entrega)
                    VALUES (?, ?, ?, ?, ?)
                    ''', (eid, ca, qtd, desc, dt))

    # 4. Scan ATESTADOS folder and link
    if os.path.exists(ATESTADOS_DIR):
        cursor.execute("SELECT id, nome FROM funcionarios WHERE status = 'Ativo'")
        active_emps = cursor.fetchall()
        for atest_file in os.listdir(ATESTADOS_DIR):
            if atest_file.lower().endswith(('.pdf', '.jpeg', '.jpg', '.png')):
                upper_f = atest_file.upper()
                for eid, ename in active_emps:
                    first_name = ename.split()[0].upper()
                    if len(first_name) >= 4 and first_name in upper_f:
                        rel_atest = os.path.join('ATESTADOS', atest_file).replace('\\', '/')
                        cursor.execute('''
                        INSERT INTO atestados (funcionario_id, data_inicio, dias_afastamento, data_retorno, motivo, anexo_path)
                        VALUES (?, ?, ?, ?, ?, ?)
                        ''', (eid, '2026-08-20', 3, '2026-08-23', 'Atestado Médico Digitalizado', rel_atest))
                        break

    conn.commit()
    conn.close()

    print(f"Sucesso! Importados {total_ativos} funcionários ATIVOS e {total_desligados} DESLIGADOS.")
    print(f"Banco de dados salvo em: {DB_PATH}")

if __name__ == '__main__':
    run_import()
